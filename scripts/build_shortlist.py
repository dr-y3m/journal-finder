#!/usr/bin/env python3
"""
build_shortlist.py: build the journal-finder Excel workbook.

Usage:
    python build_shortlist.py input.json output.xlsx [--no-recalc]

Reads the JSON described in references/data-schema.md, applies the fit gate,
orders the journals by overall score, and writes a workbook whose scores are
live formulas driven by the Settings sheet. Then runs the xlsx skill's
recalc.py so that previewers see computed values.
"""
import json
import os
import re
import subprocess
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

RECALC = "/mnt/skills/public/xlsx/scripts/recalc.py"
FONT = "Arial"

DEFAULT_WEIGHTS = {"prestige": 0.40, "fit": 0.35, "speed": 0.15, "apc": 0.10}
DEFAULT_LIST_WEIGHTS = {"FNEGE": 0.30, "AJG": 0.25, "ABDC": 0.20,
                        "CNRS": 0.10, "SJR": 0.10, "WOS": 0.05}
DEFAULT_THRESHOLDS = {"stretch": 4.25, "target": 3.0, "fit_gate": 3.0}

SCORED_LISTS = ["FNEGE", "CNRS", "AJG", "ABDC", "SJR", "WOS"]
LIST_HEADERS = {
    "FNEGE": "FNEGE 2025", "CNRS": "CNRS 2020 (v5.07)", "HCERES": "HCERES",
    "AJG": "AJG 2024", "ABDC": "ABDC 2025", "SJR": "Scopus SJR quartile",
    "WOS": "WoS quartile",
}

SCALES = (
    [("FNEGE", l, s) for l, s in [("1*", 5), ("1", 4.5), ("2", 3.5), ("3", 2.5), ("4", 1.5), ("Not listed", 0)]]
    # CNRS has no "Not listed" score on purpose: the list dates from 2020, so absence often reflects a
    # journal's age rather than its standing. A CNRS "Not listed" is shown but left out of the mean.
    + [("CNRS", l, s) for l, s in [("1*", 5), ("1", 4.5), ("2", 3.5), ("3", 2.5), ("4", 1.5)]]
    + [("AJG", l, s) for l, s in [("4*", 5), ("4", 4.5), ("3", 3.5), ("2", 2.5), ("1", 1.5), ("Not listed", 0)]]
    + [("ABDC", l, s) for l, s in [("A*", 5), ("A", 4), ("B", 3), ("C", 1.5), ("Not listed", 0)]]
    + [("SJR", l, s) for l, s in [("Q1", 4), ("Q2", 3), ("Q3", 2), ("Q4", 1), ("Not listed", 0)]]
    + [("WOS", l, s) for l, s in [("Q1", 4), ("Q2", 3), ("Q3", 2), ("Q4", 1), ("Not listed", 0)]]
)
SCALE_LOOKUP = {(lst, lab): s for lst, lab, s in SCALES}

SPEED_BRACKETS = [(30, 5), (60, 4), (90, 3), (120, 2)]
SPEED_OVER, SPEED_UNKNOWN = 1, 2.5
APC_NONE = 5
APC_BRACKETS = [(1000, 3.5), (2500, 2.5), (3500, 1.5)]
APC_OVER, APC_UNKNOWN = 0.5, 2.5

FIT_KEYS = ["scope", "conversation", "method", "context", "contribution"]

# ---------- styles ----------
HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
CALC_HDR_FILL = PatternFill("solid", fgColor="7F7F7F")
BODY = Font(name=FONT, size=10)
INPUT = Font(name=FONT, size=10, color="0000FF")
BOLD = Font(name=FONT, size=10, bold=True)
TITLE = Font(name=FONT, size=14, bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")
CONF_FILL = {"O": PatternFill("solid", fgColor="C6EFCE"),
             "S": PatternFill("solid", fgColor="FFEB9C"),
             "U": PatternFill("solid", fgColor="EDEDED")}


# ---------- normalisation ----------
def norm_label(lst, value):
    """Return (label, recognised) for a ranking value."""
    if value is None:
        return "", False
    v = str(value).strip()
    if not v:
        return "", False
    if v.lower() in {"not listed", "non classée", "non classé", "absent", "not ranked"}:
        return "Not listed", True
    if lst in ("FNEGE", "CNRS"):
        m = re.fullmatch(r"(?:cat(?:[ée]gor(?:ie|y))?\.?\s*)?([1-4])\s*(\*?)", v, re.I)
        if m and (not m.group(2) or m.group(1) == "1"):
            return m.group(1) + m.group(2), True
    elif lst == "AJG":
        m = re.fullmatch(r"(?:(?:ajg|abs)\s*)?([1-4])\s*(\*?)", v, re.I)
        if m and (not m.group(2) or m.group(1) == "4"):
            return m.group(1) + m.group(2), True
    elif lst == "ABDC":
        m = re.fullmatch(r"(A)\s*(\*?)|([BC])", v.upper())
        if m:
            return (m.group(1) + m.group(2)) if m.group(1) else m.group(3), True
    elif lst in ("SJR", "WOS"):
        m = re.fullmatch(r"q\s*([1-4])", v, re.I)
        if m:
            return "Q" + m.group(1), True
    return v, False


def num(x):
    try:
        if x is None or x == "":
            return None
        return float(x)
    except (TypeError, ValueError):
        return None


def yes_no(x):
    if x is True or (isinstance(x, str) and x.strip().lower() in {"yes", "true", "oui"}):
        return "Yes"
    if x is False or (isinstance(x, str) and x.strip().lower() in {"no", "false", "non"}):
        return "No"
    return ""


def display_ranking(j, lst):
    r = (j.get("rankings") or {}).get(lst) or {}
    conf = str(r.get("confidence") or "U").upper()[:1]
    if conf not in CONF_FILL:
        conf = "U"
    label, ok = norm_label(lst, r.get("value"))
    if lst == "HCERES":
        shown = label if (label and conf != "U") else "not verified"
        return shown, conf, r, False
    if conf == "U" or not label or label.lower() == "not verified":
        return "not verified", "U", r, False
    if label == "Not listed" and conf != "O":
        return "not verified", "U", r, False
    return label, conf, r, ok


# ---------- python mirror of the workbook formulas (used for ordering) ----------
def compute(j, w, lw):
    subs = [num((j.get("fit") or {}).get(k)) for k in FIT_KEYS]
    subs = [s for s in subs if s is not None]
    fit = round(sum(subs) / len(subs), 2) if subs else None

    tot, den = 0.0, 0.0
    for lst in SCORED_LISTS:
        shown, conf, _, ok = display_ranking(j, lst)
        if ok:
            s = SCALE_LOOKUP.get((lst, shown))
            if s is not None:
                tot += s * lw[lst]
                den += lw[lst]
    prestige = round(tot / den, 2) if den else None

    days = num(j.get("days_to_first_decision"))
    if days is None:
        speed = SPEED_UNKNOWN
    else:
        speed = next((s for lim, s in SPEED_BRACKETS if days <= lim), SPEED_OVER)

    mand, eur = yes_no(j.get("apc_mandatory")), num(j.get("apc_eur"))
    if mand == "No":
        apc = APC_NONE
    elif mand == "Yes":
        apc = APC_UNKNOWN if eur is None else next((s for lim, s in APC_BRACKETS if eur <= lim), APC_OVER)
    else:
        apc = APC_UNKNOWN

    overall = None
    if fit is not None:
        wsum = sum(w.values())
        overall = round(20 * (w["prestige"] * (prestige or 0) + w["fit"] * fit
                              + w["speed"] * speed + w["apc"] * apc) / wsum, 1)
    return {"fit": fit, "prestige": prestige, "speed": speed, "apc": apc, "overall": overall}


# ---------- sheet helpers ----------
def header_row(ws, headers, calc_from=None):
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = HDR_FONT
        cell.fill = CALC_HDR_FILL if (calc_from and c >= calc_from) else HDR_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 42
    ws.freeze_panes = "C2"


def set_widths(ws, widths):
    for c, wdt in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c)].width = wdt


def put(ws, r, c, v, font=BODY, align=TOP, link=None):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font = font
    cell.alignment = align
    if link:
        cell.hyperlink = link
    return cell


def build_settings(wb, w, lw, th):
    ws = wb.create_sheet("Settings")
    put(ws, 1, 1, "Settings: edit the blue cells; every score on the Shortlist recalculates", TITLE)
    refs = {}

    put(ws, 3, 1, "Component weight", BOLD)
    for i, k in enumerate(["prestige", "fit", "speed", "apc"]):
        r = 4 + i
        put(ws, r, 1, k.capitalize() if k != "apc" else "APC")
        put(ws, r, 2, w[k], INPUT).number_format = "0%"
        refs["w_" + k] = f"Settings!$B${r}"

    put(ws, 9, 1, "List weight within prestige", BOLD)
    for i, lst in enumerate(["FNEGE", "AJG", "ABDC", "CNRS", "SJR", "WOS"]):
        r = 10 + i
        put(ws, r, 1, LIST_HEADERS[lst])
        put(ws, r, 2, lw[lst], INPUT).number_format = "0%"
        refs["lw_" + lst] = f"Settings!$B${r}"

    put(ws, 17, 1, "Thresholds", BOLD)
    put(ws, 18, 1, "Stretch if prestige ≥")
    put(ws, 18, 2, th["stretch"], INPUT)
    put(ws, 19, 1, "Target if prestige ≥ (else Safe)")
    put(ws, 19, 2, th["target"], INPUT)
    put(ws, 20, 1, "Fit gate (applied when the file was built)")
    put(ws, 20, 2, th["fit_gate"], INPUT)
    refs["stretch"], refs["target"] = "Settings!$B$18", "Settings!$B$19"

    put(ws, 22, 1, "Speed score: days to first decision", BOLD)
    put(ws, 22, 2, "Up to (days)", BOLD)
    put(ws, 22, 3, "Score", BOLD)
    for i, (lim, s) in enumerate(SPEED_BRACKETS):
        r = 23 + i
        put(ws, r, 1, f"Bracket {i + 1}")
        put(ws, r, 2, lim, INPUT)
        put(ws, r, 3, s, INPUT)
    put(ws, 27, 1, "Longer than the last bracket")
    put(ws, 27, 3, SPEED_OVER, INPUT)
    put(ws, 28, 1, "Unknown")
    put(ws, 28, 3, SPEED_UNKNOWN, INPUT)

    put(ws, 30, 1, "APC score", BOLD)
    put(ws, 30, 2, "Up to (≈ EUR)", BOLD)
    put(ws, 30, 3, "Score", BOLD)
    put(ws, 31, 1, "No mandatory APC")
    put(ws, 31, 3, APC_NONE, INPUT)
    for i, (lim, s) in enumerate(APC_BRACKETS):
        r = 32 + i
        put(ws, r, 1, f"Mandatory APC, bracket {i + 1}")
        put(ws, r, 2, lim, INPUT)
        put(ws, r, 3, s, INPUT)
    put(ws, 35, 1, "Mandatory APC above the last bracket")
    put(ws, 35, 3, APC_OVER, INPUT)
    put(ws, 36, 1, "Unknown")
    put(ws, 36, 3, APC_UNKNOWN, INPUT)

    put(ws, 3, 5, "Ranking scales (0–5)", BOLD)
    put(ws, 4, 5, "Key", BOLD)
    put(ws, 4, 6, "List", BOLD)
    put(ws, 4, 7, "Label", BOLD)
    put(ws, 4, 8, "Score", BOLD)
    first = 5
    for i, (lst, lab, s) in enumerate(SCALES):
        r = first + i
        put(ws, r, 5, f'=F{r}&"_"&SUBSTITUTE(G{r},"*","star")')
        put(ws, r, 6, lst)
        put(ws, r, 7, lab)
        put(ws, r, 8, s, INPUT)
    last = first + len(SCALES) - 1
    refs["scale_keys"] = f"Settings!$E${first}:$E${last}"
    refs["scale_vals"] = f"Settings!$H${first}:$H${last}"

    put(ws, last + 2, 5, "Unverified lists are left out of the prestige mean; a verified 'Not listed' scores 0, except for "
                         "CNRS (2020 list: absence often reflects a journal's age, so it is left out). "
                         "HCERES is shown but not scored (derived from FNEGE and CNRS).", WRAP)
    set_widths(ws, [44, 14, 10, 3, 16, 10, 12, 8])
    return refs


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if len(args) != 2:
        print(__doc__)
        sys.exit(1)
    src, out = args
    recalc = "--no-recalc" not in sys.argv

    with open(src, encoding="utf-8") as f:
        data = json.load(f)

    w = {**DEFAULT_WEIGHTS, **(data.get("weights") or {})}
    lw = {**DEFAULT_LIST_WEIGHTS, **(data.get("list_weights") or {})}
    th = {**DEFAULT_THRESHOLDS, **(data.get("thresholds") or {})}
    ms = data.get("manuscript") or {}

    shortlist, excluded = [], list(data.get("excluded") or [])
    for j in data.get("journals") or []:
        sc = compute(j, w, lw)
        if sc["fit"] is None or sc["fit"] < th["fit_gate"]:
            fit_txt = "no fit score" if sc["fit"] is None else f"fit {sc['fit']:.2f}"
            excluded.append({"name": j.get("name"), "reason": f"Below the fit gate ({fit_txt} < {th['fit_gate']})",
                             "channels": j.get("channels"), "fit": sc["fit"]})
        elif j.get("red_flags"):
            excluded.append({"name": j.get("name"), "reason": "Red flags: " + str(j["red_flags"]),
                             "channels": j.get("channels"), "fit": sc["fit"]})
        else:
            shortlist.append((j, sc))
    shortlist.sort(key=lambda t: (t[1]["overall"] is None, -(t[1]["overall"] or 0), -(t[1]["fit"] or 0)))

    wb = Workbook()
    ws = wb.active
    ws.title = "Shortlist"
    refs = build_settings(wb, w, lw, th)

    headers = (["Rank", "Journal", "Publisher", "Tier", "Overall /100", "Prestige /5", "Fit /5",
                "Speed /5", "APC /5"]
               + [LIST_HEADERS[l] for l in ["FNEGE", "CNRS", "HCERES", "AJG", "ABDC", "SJR", "WOS"]]
               + ["Days to 1st decision", "Acceptance rate", "OA model / APC", "APC mandatory?",
                  "APC (≈ EUR)", "Length limit", "Method stance", "Fit rationale", "Desk-reject risks",
                  "Papers to engage with", "Open special issue", "Channels", "Homepage"]
               + ["Fit: scope", "Fit: conversation", "Fit: method", "Fit: context", "Fit: contribution"]
               + ["Score: FNEGE", "Score: CNRS", "Score: AJG", "Score: ABDC", "Score: SJR", "Score: WoS"])
    col = {h: i + 1 for i, h in enumerate(headers)}
    L = {h: get_column_letter(i) for h, i in col.items()}
    calc_from = col["Score: FNEGE"]
    header_row(ws, headers, calc_from=calc_from)

    rank_cols = {"FNEGE": LIST_HEADERS["FNEGE"], "CNRS": LIST_HEADERS["CNRS"], "HCERES": LIST_HEADERS["HCERES"],
                 "AJG": LIST_HEADERS["AJG"], "ABDC": LIST_HEADERS["ABDC"], "SJR": LIST_HEADERS["SJR"],
                 "WOS": LIST_HEADERS["WOS"]}
    helper_cols = {"FNEGE": "Score: FNEGE", "CNRS": "Score: CNRS", "AJG": "Score: AJG",
                   "ABDC": "Score: ABDC", "SJR": "Score: SJR", "WOS": "Score: WoS"}

    evidence = []
    for i, (j, sc) in enumerate(shortlist):
        r = i + 2
        put(ws, r, col["Rank"], i + 1)
        put(ws, r, col["Journal"], j.get("name"), BOLD, WRAP)
        put(ws, r, col["Publisher"], j.get("publisher"), BODY, WRAP)

        for lst, h in rank_cols.items():
            shown, conf, raw, _ = display_ranking(j, lst)
            cell = put(ws, r, col[h], shown, BODY, Alignment(horizontal="center", vertical="top"))
            cell.fill = CONF_FILL[conf]
            parts = [f"Confidence: {conf}"]
            for k in ("edition", "detail", "source", "checked", "note"):
                if raw.get(k):
                    parts.append(f"{k.capitalize()}: {raw[k]}")
            cell.comment = Comment("\n".join(parts), "journal-finder", width=320, height=140)
            evidence.append([j.get("name"), LIST_HEADERS[lst], shown, raw.get("edition"), conf,
                             raw.get("detail"), raw.get("source"), raw.get("checked"), raw.get("note")])

        put(ws, r, col["Days to 1st decision"], num(j.get("days_to_first_decision")), INPUT)
        put(ws, r, col["Acceptance rate"], j.get("acceptance_rate"), BODY, WRAP)
        put(ws, r, col["OA model / APC"], j.get("oa_model"), BODY, WRAP)
        put(ws, r, col["APC mandatory?"], yes_no(j.get("apc_mandatory")), INPUT)
        put(ws, r, col["APC (≈ EUR)"], num(j.get("apc_eur")), INPUT).number_format = "#,##0"
        put(ws, r, col["Length limit"], j.get("length_limit"), BODY, WRAP)
        put(ws, r, col["Method stance"], j.get("method_stance"), BODY, WRAP)
        put(ws, r, col["Fit rationale"], j.get("fit_rationale"), BODY, WRAP)
        put(ws, r, col["Desk-reject risks"], j.get("desk_reject_risks"), BODY, WRAP)
        papers = j.get("recent_papers") or []
        put(ws, r, col["Papers to engage with"],
            "\n".join(f"{p.get('title')} ({p.get('year', 'n.d.')}) {p.get('url', '')}".strip() for p in papers) or None,
            BODY, WRAP)
        si = j.get("special_issue") or {}
        put(ws, r, col["Open special issue"],
            f"{si.get('title')} (deadline {si.get('deadline', 'n/a')}) {si.get('url', '')}".strip() if si else None,
            BODY, WRAP)
        put(ws, r, col["Channels"], ", ".join(j.get("channels") or []) or None)
        hp = j.get("homepage")
        put(ws, r, col["Homepage"], hp, Font(name=FONT, size=10, color="0563C1", underline="single") if hp else BODY,
            WRAP, link=hp)
        for k in FIT_KEYS:
            put(ws, r, col["Fit: " + k], num((j.get("fit") or {}).get(k)), INPUT)

        # helper scores: map each ranking label to its 0-5 score via the Settings table
        for lst, h in helper_cols.items():
            v = f"{L[rank_cols[lst]]}{r}"
            put(ws, r, col[h],
                f'=IFERROR(INDEX({refs["scale_vals"]},MATCH("{lst}_"&SUBSTITUTE({v},"*","star"),{refs["scale_keys"]},0)),"")')

        hs = {lst: f"{L[h]}{r}" for lst, h in helper_cols.items()}
        num_terms = "+".join(f'IF({hs[l]}="",0,{hs[l]}*{refs["lw_" + l]})' for l in SCORED_LISTS)
        den_terms = "+".join(f'IF({hs[l]}="",0,{refs["lw_" + l]})' for l in SCORED_LISTS)
        P, F, S, A = (f"{L['Prestige /5']}{r}", f"{L['Fit /5']}{r}", f"{L['Speed /5']}{r}", f"{L['APC /5']}{r}")
        fit_rng = f"{L['Fit: scope']}{r}:{L['Fit: contribution']}{r}"
        D = f"{L['Days to 1st decision']}{r}"
        M, E = f"{L['APC mandatory?']}{r}", f"{L['APC (≈ EUR)']}{r}"

        put(ws, r, col["Prestige /5"], f'=IF(({den_terms})=0,"",ROUND(({num_terms})/({den_terms}),2))').number_format = "0.00"
        put(ws, r, col["Fit /5"], f'=IF(COUNT({fit_rng})=0,"",ROUND(AVERAGE({fit_rng}),2))').number_format = "0.00"
        put(ws, r, col["Speed /5"],
            f'=IF({D}="",Settings!$C$28,IF({D}<=Settings!$B$23,Settings!$C$23,IF({D}<=Settings!$B$24,Settings!$C$24,'
            f'IF({D}<=Settings!$B$25,Settings!$C$25,IF({D}<=Settings!$B$26,Settings!$C$26,Settings!$C$27)))))')
        put(ws, r, col["APC /5"],
            f'=IF({M}="No",Settings!$C$31,IF({M}="Yes",IF({E}="",Settings!$C$36,IF({E}<=Settings!$B$32,Settings!$C$32,'
            f'IF({E}<=Settings!$B$33,Settings!$C$33,IF({E}<=Settings!$B$34,Settings!$C$34,Settings!$C$35)))),Settings!$C$36))')
        wsum = f'({refs["w_prestige"]}+{refs["w_fit"]}+{refs["w_speed"]}+{refs["w_apc"]})'
        put(ws, r, col["Overall /100"],
            f'=IF({F}="","",ROUND(20*({refs["w_prestige"]}*IF({P}="",0,{P})+{refs["w_fit"]}*{F}'
            f'+{refs["w_speed"]}*{S}+{refs["w_apc"]}*{A})/{wsum},1))', BOLD).number_format = "0.0"
        put(ws, r, col["Tier"],
            f'=IF({P}="","Unverified",IF({P}>={refs["stretch"]},"Stretch",IF({P}>={refs["target"]},"Target","Safe")))', BOLD)

    last_row = max(2, len(shortlist) + 1)
    tier_rng = f"{L['Tier']}2:{L['Tier']}{last_row}"
    for label, color in [("Stretch", "E4DFEC"), ("Target", "DDEBF7"), ("Safe", "E2EFDA"), ("Unverified", "EDEDED")]:
        ws.conditional_formatting.add(tier_rng, FormulaRule(formula=[f'$' + L['Tier'] + f'2="{label}"'],
                                                            fill=PatternFill("solid", fgColor=color)))
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last_row}"
    widths = ([6, 30, 16, 10, 9, 9, 8, 8, 8] + [10, 11, 14, 9, 9, 11, 10]
              + [10, 12, 18, 10, 10, 16, 30, 40, 36, 50, 34, 10, 30] + [9] * 5 + [9] * 6)
    set_widths(ws, widths)

    # Ranking evidence
    ev = wb.create_sheet("Ranking evidence", 1)
    ev_headers = ["Journal", "List", "Value shown", "Edition", "Confidence", "Detail", "Source", "Checked", "Note"]
    header_row(ev, ev_headers)
    for i, row in enumerate(evidence):
        for c, v in enumerate(row, start=1):
            cell = put(ev, i + 2, c, v, BODY, WRAP, link=v if (c == 7 and v and str(v).startswith("http")) else None)
            if c == 5:
                cell.fill = CONF_FILL.get(v, CONF_FILL["U"])
    set_widths(ev, [30, 18, 14, 14, 11, 34, 50, 12, 36])

    # Special issues
    si_ws = wb.create_sheet("Special issues", 2)
    header_row(si_ws, ["Journal", "Special issue", "Deadline", "URL", "Note"])
    sis, seen = [], set()
    for s in data.get("special_issues") or []:
        sis.append(s)
    for j, _ in shortlist:
        s = j.get("special_issue")
        if s:
            sis.append({"journal": j.get("name"), **s})
    for i, s in enumerate(x for x in sis if (x.get("journal"), x.get("title")) not in seen and not seen.add((x.get("journal"), x.get("title")))):
        for c, k in enumerate(["journal", "title", "deadline", "url", "note"], start=1):
            v = s.get(k)
            put(si_ws, i + 2, c, v, BODY, WRAP, link=v if (k == "url" and v) else None)
    set_widths(si_ws, [30, 50, 12, 50, 40])

    # Excluded
    ex = wb.create_sheet("Excluded", 3)
    header_row(ex, ["Journal", "Reason", "Fit /5", "Channels"])
    for i, e in enumerate(excluded):
        put(ex, i + 2, 1, e.get("name"), BOLD, WRAP)
        put(ex, i + 2, 2, e.get("reason"), BODY, WRAP)
        put(ex, i + 2, 3, e.get("fit"))
        put(ex, i + 2, 4, ", ".join(e.get("channels") or []) or None)
    set_widths(ex, [34, 70, 8, 12])

    # Fingerprint
    fp = wb.create_sheet("Fingerprint", 4)
    header_row(fp, ["Field", "Value"])
    rows = [("Title", ms.get("title")), ("Input mode", ms.get("input_mode")),
            ("Journal language scope", ms.get("language_scope")), ("Constraints", ms.get("constraints"))]
    rows += list((ms.get("fingerprint") or {}).items())
    for i, (k, v) in enumerate(rows):
        put(fp, i + 2, 1, k, BOLD, WRAP)
        put(fp, i + 2, 2, v, BODY, WRAP)
    set_widths(fp, [26, 100])

    # About
    ab = wb.create_sheet("About")
    lines = [
        ("Journal shortlist", TITLE),
        (f"Manuscript: {ms.get('title') or 'n/a'}", BOLD),
        (f"Built on {ms.get('date') or date.today().isoformat()} by the journal-finder skill.", BODY),
        ("", BODY),
        ("How to read the Shortlist", BOLD),
        ("Rows are ordered by Overall score at build time. After changing weights on the Settings sheet, "
         "re-sort with Data > Sort on the Overall column.", BODY),
        ("Overall = weighted mean of Prestige, Fit, Speed and APC (each 0–5), scaled to 100.", BODY),
        ("Prestige = weighted mean of the verified rankings mapped to 0–5 (Settings sheet). Unverified lists are "
         "left out; a verified 'Not listed' counts as 0 (except CNRS, whose 2020 list is left out when it doesn't list the journal). HCERES is shown but not scored.", BODY),
        ("Tier: Stretch / Target / Safe from Prestige thresholds. Only journals that passed the fit gate appear; "
         "the others are on the Excluded sheet.", BODY),
        ("", BODY),
        ("Ranking cell colours (hover a cell for edition and source)", BOLD),
        ("Green = official source consulted (O) · Amber = explicit secondary source naming the edition (S) · "
         "Grey = not verified (U).", BODY),
        ("Blue text = inputs you can edit; black = formulas.", BODY),
    ]
    for i, (t, f) in enumerate(lines):
        put(ab, i + 1, 1, t, f, WRAP)
    ab.column_dimensions["A"].width = 120

    wb.save(out)
    print(f"Saved {out}: {len(shortlist)} shortlisted, {len(excluded)} excluded.")
    for k, (j, sc) in enumerate(shortlist[:8]):
        print(f"  {k + 1}. {j.get('name')}: overall {sc['overall']}, prestige {sc['prestige']}, fit {sc['fit']}")

    if recalc:
        if os.path.exists(RECALC):
            res = subprocess.run([sys.executable, RECALC, os.path.abspath(out), "90"],
                                 cwd=os.path.dirname(os.path.dirname(RECALC)), capture_output=True, text=True)
            print("recalc:", (res.stdout or res.stderr).strip())
        else:
            print("WARNING: recalc.py not found; formulas have no cached values until opened in Excel.")


if __name__ == "__main__":
    main()
